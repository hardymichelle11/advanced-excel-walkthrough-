import datetime as dt
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.chart import BarChart, LineChart, PieChart, ScatterChart, Reference, Series

GREEN = '217346'
HDR_FILL = PatternFill('solid', start_color=GREEN)
IN_FILL = PatternFill('solid', start_color='FFF2CC')
CHECK_FILL = PatternFill('solid', start_color='E2EFDA')
TODO_FILL = PatternFill('solid', start_color='FFFF00')
thin = Side(style='thin', color='BFBFBF')
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
USD = '$#,##0'
USD2 = '$#,##0.00'
DATE = 'mmm d, yyyy'


def font(size=10, **k):
    return Font(name='Arial', size=size, **k)


wb = Workbook()
wb.remove(wb.active)


def sheet(name, title, sub, widths):
    ws = wb.create_sheet(name)
    ws['A1'] = title
    ws['A1'].font = font(size=14, bold=True, color=GREEN)
    ws['A2'] = sub
    ws['A2'].font = font(italic=True, color='595959')
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[L(i)].width = w
    return ws


def put(ws, ref, v, kind='text', fmt=None, bold=False, align=None):
    c = ws[ref]
    c.value = v
    color = '000000'
    italic = False
    if kind == 'input':
        color = '0000FF'; c.fill = IN_FILL; c.border = BORDER
    elif kind == 'link':
        color = '008000'
    elif kind == 'ftext':
        c.data_type = 's'; color = '7F3F00'
    elif kind == 'hdr':
        c.fill = HDR_FILL; color = 'FFFFFF'; bold = True
        c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    elif kind == 'check':
        c.fill = CHECK_FILL; c.border = BORDER
    elif kind == 'todo':
        c.fill = TODO_FILL; c.border = BORDER
    elif kind == 'note':
        color = '595959'; italic = True
    if kind == 'code':
        c.font = Font(name='Courier New', size=10, color='1F3864')
    else:
        c.font = font(bold=bold, color=color, italic=italic)
    if fmt:
        c.number_format = fmt
    if align:
        c.alignment = Alignment(horizontal=align)
    return c


def hdr(ws, row, col, labels):
    for i, l in enumerate(labels):
        put(ws, f'{L(col + i)}{row}', l, 'hdr')


def steps(ws, row, lines, col=1, head=None):
    if head:
        put(ws, f'{L(col)}{row}', head, bold=True)
        row += 1
    for ln in lines:
        put(ws, f'{L(col)}{row}', ln)
        row += 1
    return row


# ---------------------------------------------------------------- Start Here
ws = sheet('Start Here', 'Advanced Excel Practice Workbook',
           'Companion to the Advanced Excel Walkthrough Guide. Each tab matches a guide topic with the same numbers.',
           [18, 34, 70])
r = steps(ws, 4, [
    'Blue text on pale yellow = an input you can change.',
    'Black text = a formula. Green text = a formula that pulls from another sheet.',
    'Bright yellow cells = empty on purpose: you build that part (the guide shows how).',
    'Pale green cells = answer checks, so you can confirm your result.',
    'Brown text = a formula shown as text for you to read or type yourself.',
], head='Colour legend')
r += 1
hdr(ws, r, 1, ['Sheet', 'Guide topic', 'What you practise'])
ws.row_dimensions[r].height = 18
index = [
    ('Data', 'Shared sales data', '20 orders used by most other sheets; price and category are looked up'),
    ('1 Charts', '1. Charting', 'Column, pie and combo charts built from summary formulas'),
    ('2 Dates', '2. Date and time functions', 'TODAY, DATE, EDATE, EOMONTH, NETWORKDAYS, DATEDIF, time maths'),
    ('3 PivotTables', '3. PivotTables', 'Build a PivotTable from Data and check it against SUMIFS'),
    ('4 External Data', '4. Internal and external data', 'Clean imported text, import the CSV, Tables and structured references'),
    ('5 Sort Filter', '5. Sorting and filtering', 'Multi-level sort, AutoFilter, Advanced Filter with a criteria range'),
    ('6 Lookup Query', '6. Extracting and querying', 'VLOOKUP, INDEX/MATCH, two-way lookup, DSUM and friends'),
    ('7 Multi-Sheet', '7. Multiple worksheets and workbooks', '3-D references across Jan, Feb, Mar and external link syntax'),
    ('8 Auditing', '8. Formula auditing', 'Trace arrows, error types and fixes, a bug hunt'),
    ('9 Input Control', '9. Controlling user input', 'Data validation rules, input messages, unlocked cells, protection'),
    ('10 Financial', '10. Financial functions', 'PMT, IPMT, PPMT, FV, PV, NPV, IRR, RATE, NPER, SLN, amortization'),
    ('11 Solver Linear', '11. Solver: linear optimization', 'Product-mix model ready for Simplex LP'),
    ('12 Solver Nonlinear', '12. Solver: nonlinear optimization', 'Price-setting model ready for GRG Nonlinear'),
    ('13 Data Tables', '13. One- and two-variable data tables', 'Loan payment sensitivity with answer checks'),
    ('14 Scenarios', '14. Scenario management', 'Best / Base / Worst scenarios and Goal Seek'),
    ('15 Macros', '15. Macros', 'Record a macro and paste three short VBA procedures'),
]
for name, topic, what in index:
    r += 1
    c = put(ws, f'A{r}', name)
    c.hyperlink = f"#'{name}'!A1"
    c.font = font(color='0563C1', underline='single')
    put(ws, f'B{r}', topic)
    put(ws, f'C{r}', what)
r += 2
steps(ws, r, [
    'Example data is invented for teaching (Jan-May 2026 orders for a small office-supply seller).',
    'Macros cannot live in an .xlsx file. On the Macros tab you save a copy as .xlsm before recording.',
    'PivotTables, data tables, scenarios and Solver runs are left for you to perform; every one has an answer check.',
], head='Notes')

# ---------------------------------------------------------------- Data
ws = sheet('Data', 'Sales data (shared by the other sheets)',
           'An Excel Table named SalesData. Category and Unit Price are looked up from the product list on the right.',
           [14, 10, 10, 12, 13, 8, 11, 12, 3, 3, 12, 13, 11])
hdr(ws, 4, 1, ['Date', 'Region', 'Rep', 'Product', 'Category', 'Units', 'Unit Price', 'Revenue'])
hdr(ws, 4, 11, ['Product', 'Category', 'Unit Price'])
products = [('Laptop', 'Hardware', 900), ('Monitor', 'Hardware', 250), ('Keyboard', 'Accessories', 40),
            ('Mouse', 'Accessories', 25), ('License', 'Software', 120)]
for i, (p, cat, price) in enumerate(products):
    put(ws, f'K{5 + i}', p, 'input')
    put(ws, f'L{5 + i}', cat, 'input')
    put(ws, f'M{5 + i}', price, 'input', USD)
put(ws, 'K11', 'Product list: invented prices. Change one and every sheet updates.', 'note')
orders = [
    ((2026, 1, 5), 'East', 'Avery', 'Laptop', 4), ((2026, 1, 12), 'West', 'Blake', 'Monitor', 10),
    ((2026, 1, 19), 'North', 'Casey', 'Keyboard', 25), ((2026, 1, 26), 'South', 'Devon', 'Mouse', 40),
    ((2026, 2, 2), 'East', 'Avery', 'License', 15), ((2026, 2, 9), 'West', 'Blake', 'Laptop', 6),
    ((2026, 2, 16), 'North', 'Casey', 'Monitor', 8), ((2026, 2, 23), 'South', 'Devon', 'Keyboard', 30),
    ((2026, 3, 2), 'East', 'Avery', 'Mouse', 50), ((2026, 3, 9), 'West', 'Blake', 'License', 20),
    ((2026, 3, 16), 'North', 'Casey', 'Laptop', 3), ((2026, 3, 23), 'South', 'Devon', 'Monitor', 12),
    ((2026, 4, 6), 'East', 'Avery', 'Keyboard', 20), ((2026, 4, 13), 'West', 'Blake', 'Mouse', 35),
    ((2026, 4, 20), 'North', 'Casey', 'License', 10), ((2026, 4, 27), 'South', 'Devon', 'Laptop', 5),
    ((2026, 5, 4), 'East', 'Avery', 'Monitor', 9), ((2026, 5, 11), 'West', 'Blake', 'Keyboard', 15),
    ((2026, 5, 18), 'North', 'Casey', 'Mouse', 60), ((2026, 5, 25), 'South', 'Devon', 'License', 25),
]
price_of = {p: pr for p, _, pr in products}
for i, (d, reg, rep, prod, units) in enumerate(orders):
    r = 5 + i
    put(ws, f'A{r}', dt.date(*d), 'input', DATE)
    put(ws, f'B{r}', reg, 'input')
    put(ws, f'C{r}', rep, 'input')
    put(ws, f'D{r}', prod, 'input')
    put(ws, f'E{r}', f'=VLOOKUP(D{r},$K$5:$M$9,2,FALSE)')
    put(ws, f'F{r}', units, 'input')
    put(ws, f'G{r}', f'=VLOOKUP(D{r},$K$5:$M$9,3,FALSE)', fmt=USD)
    put(ws, f'H{r}', f'=F{r}*G{r}', fmt=USD)
put(ws, 'E26', 'Total', bold=True)
put(ws, 'F26', '=SUM(F5:F24)', bold=True)
put(ws, 'H26', '=SUM(H5:H24)', bold=True, fmt=USD)
t = Table(displayName='SalesData', ref='A4:H24')
t.tableStyleInfo = TableStyleInfo(name='TableStyleMedium7', showRowStripes=True)
ws.add_table(t)
ws.freeze_panes = 'A5'

D = "Data!"
REV, REG, CAT, DTE, UNITS, REP = (D + '$H$5:$H$24', D + '$B$5:$B$24', D + '$E$5:$E$24',
                                  D + '$A$5:$A$24', D + '$F$5:$F$24', D + '$C$5:$C$24')

# ---------------------------------------------------------------- 1 Charts
ws = sheet('1 Charts', '1. Charting', 'Guide topic 1. Charts read from the summary tables, which read from Data.',
           [14, 9, 12, 9, 3, 12, 12, 12])
hdr(ws, 4, 1, ['Month start', 'Label', 'Revenue', 'Units'])
for i in range(5):
    r = 5 + i
    put(ws, f'A{r}', dt.date(2026, 1 + i, 1), 'input', DATE)
    put(ws, f'B{r}', f'=TEXT(A{r},"mmm")')
    put(ws, f'C{r}', f'=SUMIFS({REV},{DTE},">="&A{r},{DTE},"<="&EOMONTH(A{r},0))', 'link', USD)
    put(ws, f'D{r}', f'=SUMIFS({UNITS},{DTE},">="&A{r},{DTE},"<="&EOMONTH(A{r},0))', 'link')
put(ws, 'B10', 'Total', bold=True)
put(ws, 'C10', '=SUM(C5:C9)', bold=True, fmt=USD)
put(ws, 'D10', '=SUM(D5:D9)', bold=True)
hdr(ws, 13, 1, ['Region', 'Revenue', 'Share'])
for i, reg in enumerate(['East', 'West', 'North', 'South']):
    r = 14 + i
    put(ws, f'A{r}', reg)
    put(ws, f'B{r}', f'=SUMIFS({REV},{REG},A{r})', 'link', USD)
    put(ws, f'C{r}', f'=B{r}/$B$18', fmt='0.0%')
put(ws, 'A18', 'Total', bold=True)
put(ws, 'B18', '=SUM(B14:B17)', bold=True, fmt=USD)
steps(ws, 21, [
    '1. Click the column chart, then Chart Design > Add Chart Element > Data Labels > Outside End.',
    '2. Right-click the pie > Format Data Labels > tick Percentage.',
    '3. Rebuild the combo yourself: select B4:D9 > Insert > Recommended Charts > All Charts > Combo,',
    '   set Units to Line and tick Secondary Axis.',
    '4. Change a Units value on the Data sheet and watch all three charts move.',
], head='Try it')

cats = Reference(ws, min_col=2, min_row=5, max_row=9)
c1 = BarChart(); c1.type = 'col'; c1.title = 'Revenue by month'
c1.add_data(Reference(ws, min_col=3, min_row=4, max_row=9), titles_from_data=True)
c1.set_categories(cats); c1.legend = None
c1.y_axis.title = 'Revenue ($)'; c1.x_axis.delete = False; c1.y_axis.delete = False
c1.height = 7.5; c1.width = 13
ws.add_chart(c1, 'F4')

c2 = PieChart(); c2.title = 'Revenue share by region'
c2.add_data(Reference(ws, min_col=2, min_row=13, max_row=17), titles_from_data=True)
c2.set_categories(Reference(ws, min_col=1, min_row=14, max_row=17))
c2.height = 7.5; c2.width = 13
ws.add_chart(c2, 'F20')

c3 = BarChart(); c3.type = 'col'; c3.title = 'Combo: revenue (columns) and units (line)'
c3.add_data(Reference(ws, min_col=3, min_row=4, max_row=9), titles_from_data=True)
c3.set_categories(cats); c3.y_axis.title = 'Revenue ($)'
c3.x_axis.delete = False; c3.y_axis.delete = False
ln = LineChart()
ln.add_data(Reference(ws, min_col=4, min_row=4, max_row=9), titles_from_data=True)
ln.y_axis.axId = 200; ln.y_axis.title = 'Units'; ln.y_axis.crosses = 'max'; ln.y_axis.delete = False
c3 += ln
c3.height = 7.5; c3.width = 13
ws.add_chart(c3, 'F36')

# ---------------------------------------------------------------- 2 Dates
ws = sheet('2 Dates', '2. Date and time functions',
           'Guide topic 2. Excel stores a date as a serial number (days since Jan 0, 1900) and a time as a fraction of a day.',
           [34, 30, 20, 60])
put(ws, 'A4', 'Start date'); put(ws, 'B4', dt.date(2026, 1, 15), 'input', DATE)
put(ws, 'A5', 'End date'); put(ws, 'B5', dt.date(2026, 10, 8), 'input', DATE)
put(ws, 'A6', 'Clock in'); put(ws, 'B6', dt.time(8, 30), 'input', 'h:mm AM/PM')
put(ws, 'A7', 'Clock out'); put(ws, 'B7', dt.time(17, 15), 'input', 'h:mm AM/PM')
hdr(ws, 9, 1, ['What you want', 'Formula', 'Result', 'How it works'])
rows = [
    ("Today's date", '=TODAY()', DATE, 'Volatile: updates every time the workbook recalculates'),
    ('Date and time right now', '=NOW()', 'mmm d, yyyy h:mm AM/PM', 'Volatile; the decimal part is the time of day'),
    ('Build a date from parts', '=DATE(2026,12,25)', DATE, 'Year, month, day. Safer than typing text dates'),
    ('Year of start date', '=YEAR(B4)', '0', 'MONTH and DAY work the same way'),
    ('Month of start date', '=MONTH(B4)', '0', ''),
    ('Day of start date', '=DAY(B4)', '0', ''),
    ('Serial number behind the start date', '=B4', '0', 'Same cell, General format: the number Excel really stores'),
    ('Days between the two dates', '=B5-B4', '0', 'Dates are numbers, so you can subtract them'),
    ('Whole months between', '=DATEDIF(B4,B5,"m")', '0', 'Hidden function: no tooltip appears. "y" years, "d" days'),
    ('Fraction of a year between', '=YEARFRAC(B4,B5)', '0.000', 'Used for interest and prorating'),
    ('Same day 3 months later', '=EDATE(B4,3)', DATE, 'Negative months go backwards'),
    ('Last day of the start month', '=EOMONTH(B4,0)', DATE, '0 = this month, 1 = next month'),
    ('Day-of-week number', '=WEEKDAY(B4)', '0', '1 = Sunday ... 7 = Saturday by default'),
    ('Day name', '=TEXT(B4,"dddd")', None, 'TEXT turns a date into formatted text'),
    ('Week number of the year', '=WEEKNUM(B4)', '0', ''),
    ('Working days between (inclusive)', '=NETWORKDAYS(B4,B5)', '0', 'Skips Sat/Sun; add a holiday range as a 3rd argument'),
    ('Date 10 working days after start', '=WORKDAY(B4,10)', DATE, 'Good for due dates'),
    ('Hours worked', '=(B7-B6)*24', '0.00', 'Time is a fraction of a day, so multiply by 24 for hours'),
    ('Hour of clock-in', '=HOUR(B6)', '0', 'MINUTE and SECOND work the same way'),
    ('Build a time from parts', '=TIME(14,30,0)', 'h:mm AM/PM', 'Hour, minute, second'),
]
for i, (what, f, fmt, how) in enumerate(rows):
    r = 10 + i
    put(ws, f'A{r}', what)
    put(ws, f'B{r}', f, 'ftext')
    put(ws, f'C{r}', f, fmt=fmt, align='right')
    put(ws, f'D{r}', how, 'note')

# ---------------------------------------------------------------- 3 PivotTables
ws = sheet('3 PivotTables', '3. PivotTables',
           'Guide topic 3. Build the PivotTable yourself, then compare it with the answer checks below.',
           [16, 14, 14, 14, 14])
steps(ws, 4, [
    '1. Go to the Data sheet and click any cell inside the table.',
    '2. Insert > PivotTable > New Worksheet > OK.',
    '3. Drag Region to Rows, Category to Columns, Revenue to Values.',
    '4. Right-click a number > Number Format > Currency, 0 decimals.',
    '5. Your grand total should be the same as E21 below.',
    '6. Extra: drag Date to Rows above Region (Excel groups it by month); add a Slicer for Rep',
    '   (PivotTable Analyze > Insert Slicer); double-click a number to drill down to its rows.',
    '7. Change a Units value on Data, then PivotTable Analyze > Refresh. Pivots do not update on their own.',
], head='Build it')
put(ws, 'A15', 'Answer check 1: Sum of Revenue, Region by Category', bold=True)
hdr(ws, 16, 1, ['Region', 'Hardware', 'Accessories', 'Software', 'Grand Total'])
for i, reg in enumerate(['East', 'West', 'North', 'South']):
    r = 17 + i
    put(ws, f'A{r}', reg, 'check')
    for col in 'BCD':
        put(ws, f'{col}{r}', f'=SUMIFS({REV},{REG},$A{r},{CAT},{col}$16)', 'check', USD)
    put(ws, f'E{r}', f'=SUM(B{r}:D{r})', 'check', USD, bold=True)
put(ws, 'A21', 'Grand Total', 'check', bold=True)
for col in 'BCDE':
    put(ws, f'{col}21', f'=SUM({col}17:{col}20)', 'check', USD, bold=True)
put(ws, 'A24', 'Answer check 2: Rep in Rows; Revenue in Values three times (Count, Sum, Average)', bold=True)
hdr(ws, 25, 1, ['Rep', 'Orders', 'Total revenue', 'Average order'])
for i, rep in enumerate(['Avery', 'Blake', 'Casey', 'Devon']):
    r = 26 + i
    put(ws, f'A{r}', rep, 'check')
    put(ws, f'B{r}', f'=COUNTIFS({REP},A{r})', 'check')
    put(ws, f'C{r}', f'=SUMIFS({REV},{REP},A{r})', 'check', USD)
    put(ws, f'D{r}', f'=AVERAGEIFS({REV},{REP},A{r})', 'check', USD)

# ---------------------------------------------------------------- 4 External Data
ws = sheet('4 External Data', '4. Working with internal and external data',
           'Guide topic 4. Part A cleans text that arrived from another system. Part B brings in the companion CSV file.',
           [30, 14, 12, 12, 44])
put(ws, 'A4', 'Part A: clean imported text with formulas', bold=True)
hdr(ws, 5, 1, ['Raw text as imported', 'Name', 'Region', 'Amount', 'Technique'])
raw = ['  avery|EAST|3,600 ', 'BLAKE |west| 2,500', ' casey|North|1,000  ', 'devon |SOUTH|1,000', '  Avery|east| 1,800 ']
for i, s in enumerate(raw):
    r = 6 + i
    put(ws, f'A{r}', s, 'input')
    put(ws, f'B{r}', f'=PROPER(TRIM(LEFT(A{r},FIND("|",A{r})-1)))')
    put(ws, f'C{r}', f'=PROPER(TRIM(MID(A{r},FIND("|",A{r})+1,FIND("|",A{r},FIND("|",A{r})+1)-FIND("|",A{r})-1)))')
    put(ws, f'D{r}', f'=VALUE(SUBSTITUTE(TRIM(MID(A{r},FIND("|",A{r},FIND("|",A{r})+1)+1,20)),",",""))', fmt=USD)
put(ws, 'E6', 'LEFT up to the first | then TRIM and PROPER', 'note')
put(ws, 'E7', 'MID between the two | characters', 'note')
put(ws, 'E8', 'MID after the second |, strip commas, VALUE to number', 'note')
put(ws, 'C11', 'Total', bold=True)
put(ws, 'D11', '=SUM(D6:D10)', bold=True, fmt=USD)
r = steps(ws, 13, [
    'Faster for one-off jobs: select A6:A10 > Data > Text to Columns > Delimited > Other: |  (work on a copy first).',
], head='Text to Columns')
r = steps(ws, r + 1, [
    '1. Save regional_targets.csv (sent with this workbook) to your computer.',
    '2. Data > Get Data > From File > From Text/CSV > choose the file.',
    '3. Check the preview (comma delimiter, Target detected as a number) > Load.',
    '4. Excel adds a new sheet with a Table linked to the file. This is a query (Power Query).',
    '5. Open the CSV in Notepad, change East from 7000 to 7500, save, then Data > Refresh All.',
    '6. Compare each target with Q1 revenue on the 7 Multi-Sheet tab.',
], head='Part B: import the CSV as a refreshable query')
r = steps(ws, r + 1, [
    'The Data sheet is already a Table named SalesData (Ctrl+T makes one). Type these in any empty cell:',
], head='Part C: internal data in a Table (structured references)')
for f, why in [('=SUM(SalesData[Revenue])', 'Whole column by name; grows when rows are added'),
               ('=AVERAGE(SalesData[Units])', 'No cell addresses to maintain'),
               ('=SUMIFS(SalesData[Revenue],SalesData[Region],"West")', 'Readable criteria')]:
    put(ws, f'A{r}', f, 'ftext')
    put(ws, f'E{r}', why, 'note')
    r += 1

# ---------------------------------------------------------------- 5 Sort Filter
ws = sheet('5 Sort Filter', '5. Sorting and filtering',
           'Guide topic 5. The list below is a static copy of Data so you can sort and filter freely without affecting other sheets.',
           [14, 10, 10, 12, 8, 12, 3, 12, 12, 12, 12, 8, 12])
put(ws, 'A4', 'Practice list (values only)', bold=True)
hdr(ws, 5, 1, ['Date', 'Region', 'Rep', 'Product', 'Units', 'Revenue'])
for i, (d, reg, rep, prod, units) in enumerate(orders):
    r = 6 + i
    put(ws, f'A{r}', dt.date(*d), fmt=DATE)
    put(ws, f'B{r}', reg); put(ws, f'C{r}', rep); put(ws, f'D{r}', prod)
    put(ws, f'E{r}', units); put(ws, f'F{r}', units * price_of[prod], fmt=USD)
put(ws, 'H4', 'Criteria range (for Advanced Filter)', bold=True)
hdr(ws, 5, 8, ['Region', 'Revenue'])
put(ws, 'H6', 'West', 'input')
put(ws, 'I6', '>2000', 'input')
put(ws, 'H7', 'Same row = AND. A second criteria row would mean OR.', 'note')
r = steps(ws, 9, [
    '1. Click any cell in A5:F25.',
    '2. Data > Advanced.',
    '3. List range: $A$5:$F$25',
    '4. Criteria range: $H$5:$I$6',
    '5. Choose "Copy to another location", Copy to: $H$22',
    '6. OK. You should get the 3 rows shown in the check.',
], col=8, head='Advanced Filter steps')
put(ws, f'H{r + 1}', 'Answer check: West orders over $2,000', bold=True)
hdr(ws, r + 2, 8, ['Date', 'Product', 'Revenue'])
rr = r + 3
for (d, reg, rep, prod, units) in orders:
    if reg == 'West' and units * price_of[prod] > 2000:
        put(ws, f'H{rr}', dt.date(*d), 'check', DATE)
        put(ws, f'I{rr}', prod, 'check')
        put(ws, f'J{rr}', units * price_of[prod], 'check', USD)
        rr += 1
put(ws, 'H22', 'Advanced Filter output lands here', 'todo')
r = steps(ws, 28, [
    'Sort: Data > Sort > Sort by Region (A to Z) > Add Level > then by Revenue (Largest to Smallest).',
    'AutoFilter: Data > Filter, then use the Product arrow > untick Select All > tick Laptop.',
    'Number filter: Revenue arrow > Number Filters > Top 10 > change to Top 5.',
    'Clear: Data > Clear. Remove arrows: Data > Filter again.',
], head='Sort and AutoFilter')
r = steps(ws, r + 1, ['Type these in an empty area to the right. Each one spills into as many cells as it needs.'],
          head='Dynamic array versions (Excel 365 / 2021)')
for f, why in [('=SORT(A6:F25,6,-1)', 'Sort by the 6th column, descending'),
               ('=FILTER(A6:F25,(B6:B25="West")*(F6:F25>2000))', '* means AND; the same 3 rows as the check'),
               ('=UNIQUE(B6:B25)', 'The distinct regions')]:
    put(ws, f'A{r}', f, 'ftext')
    put(ws, f'H{r}', why, 'note')
    r += 1

# ---------------------------------------------------------------- 6 Lookup Query
ws = sheet('6 Lookup Query', '6. Extracting and querying data',
           'Guide topic 6. Lookups pull one matching value; database functions summarise every row that meets a criteria range.',
           [30, 16, 3, 3, 16, 14, 3, 44])
put(ws, 'A4', 'Product to look up'); put(ws, 'B4', 'Monitor', 'input')
dv = DataValidation(type='list', formula1='=Data!$K$5:$K$9', allow_blank=False)
ws.add_data_validation(dv); dv.add('B4')
put(ws, 'A5', 'Price with VLOOKUP (exact)')
put(ws, 'B5', '=VLOOKUP(B4,Data!$K$5:$M$9,3,FALSE)', 'link', USD)
put(ws, 'H5', '=VLOOKUP(B4,Data!$K$5:$M$9,3,FALSE)', 'ftext')
put(ws, 'A6', 'Category with INDEX/MATCH')
put(ws, 'B6', '=INDEX(Data!$L$5:$L$9,MATCH(B4,Data!$K$5:$K$9,0))', 'link')
put(ws, 'H6', '=INDEX(Data!$L$5:$L$9,MATCH(B4,Data!$K$5:$K$9,0))', 'ftext')
put(ws, 'A7', 'Price with XLOOKUP (type it yourself)')
put(ws, 'B7', None, 'todo')
put(ws, 'H7', '=XLOOKUP(B4,Data!K5:K9,Data!M5:M9,"Not found")', 'ftext')

put(ws, 'E4', 'Commission tiers', bold=True)
hdr(ws, 5, 5, ['Revenue from', 'Rate'])
for i, (lo, rate) in enumerate([(0, 0.02), (5000, 0.04), (10000, 0.06)]):
    put(ws, f'E{6 + i}', lo, 'input', USD)
    put(ws, f'F{6 + i}', rate, 'input', '0%')
put(ws, 'E9', 'Invented tiers; must be sorted ascending', 'note')

put(ws, 'A10', 'Rep revenue'); put(ws, 'B10', 9700, 'input', USD)
put(ws, 'A11', 'Rate with VLOOKUP (approximate)')
put(ws, 'B11', '=VLOOKUP(B10,$E$6:$F$8,2,TRUE)', fmt='0%')
put(ws, 'H11', '=VLOOKUP(B10,$E$6:$F$8,2,TRUE)', 'ftext')
put(ws, 'A12', 'Commission'); put(ws, 'B12', '=B10*B11', fmt=USD)

put(ws, 'A14', 'Region'); put(ws, 'B14', 'West', 'input')
put(ws, 'A15', 'Category'); put(ws, 'B15', 'Hardware', 'input')
dv1 = DataValidation(type='list', formula1='"East,West,North,South"'); ws.add_data_validation(dv1); dv1.add('B14')
dv2 = DataValidation(type='list', formula1='"Hardware,Accessories,Software"'); ws.add_data_validation(dv2); dv2.add('B15')
two = ("=INDEX('3 PivotTables'!$B$17:$D$20,MATCH(B14,'3 PivotTables'!$A$17:$A$20,0),"
       "MATCH(B15,'3 PivotTables'!$B$16:$D$16,0))")
put(ws, 'A16', 'Two-way lookup: INDEX/MATCH/MATCH')
put(ws, 'B16', two, 'link', USD)
put(ws, 'H16', two, 'ftext')

put(ws, 'A19', 'Database functions', bold=True)
put(ws, 'E19', 'Criteria range', bold=True)
hdr(ws, 20, 5, ['Region', 'Revenue'])
put(ws, 'E21', 'West', 'input'); put(ws, 'F21', '>2000', 'input')
db = 'Data!$A$4:$H$24'
for i, (label, fn, fmt) in enumerate([('Total revenue of matching orders', 'DSUM', USD),
                                      ('Number of matching orders', 'DCOUNT', '0'),
                                      ('Average matching order', 'DAVERAGE', USD2),
                                      ('Largest matching order', 'DMAX', USD)]):
    r = 20 + i
    f = f'={fn}({db},"Revenue",$E$20:$F$21)'
    put(ws, f'A{r}', label)
    put(ws, f'B{r}', f, 'link', fmt)
    put(ws, f'H{r}', f, 'ftext')
put(ws, 'A24', 'Same total with SUMIFS')
f = f'=SUMIFS({REV},{REG},E21,{REV},F21)'
put(ws, 'B24', f, 'link', USD)
put(ws, 'H24', f, 'ftext')
steps(ws, 27, [
    'Change B4 to Laptop: price becomes $900 and category stays Hardware.',
    'Change B10 to 10000: the rate jumps to 6% because approximate match takes the last tier that is not above the value.',
    'Change E21 to South: all four database results update at once.',
], head='Try it')

# ---------------------------------------------------------------- Jan Feb Mar + 7 Multi-Sheet
month_data = {
    'Jan': [(3600, 2200), (2500, 1600), (1000, 700), (1000, 650)],
    'Feb': [(1800, 1100), (5400, 3300), (2000, 1300), (1200, 800)],
    'Mar': [(1250, 800), (2400, 1500), (2700, 1700), (3000, 1850)],
}
regions = ['East', 'West', 'North', 'South']
ws7 = sheet('7 Multi-Sheet', '7. Models with multiple worksheets and workbooks',
            'Guide topic 7. Jan, Feb and Mar share one layout, so a 3-D reference can add the same cell across all three.',
            [14, 12, 12, 12, 16, 16, 3, 44])
for m, vals in month_data.items():
    ws = sheet(m, f'{m} 2026 results', 'Identical layout on Jan, Feb and Mar. Revenue matches the Data sheet; expenses are invented.',
               [14, 12, 12, 12])
    hdr(ws, 4, 1, ['Region', 'Revenue', 'Expenses', 'Profit'])
    for i, (reg, (rev, exp)) in enumerate(zip(regions, vals)):
        r = 5 + i
        put(ws, f'A{r}', reg)
        put(ws, f'B{r}', rev, 'input', USD)
        put(ws, f'C{r}', exp, 'input', USD)
        put(ws, f'D{r}', f'=B{r}-C{r}', fmt=USD)
    put(ws, 'A9', 'Total', bold=True)
    for col in 'BCD':
        put(ws, f'{col}9', f'=SUM({col}5:{col}8)', bold=True, fmt=USD)
ws = ws7
hdr(ws, 4, 1, ['Region', 'Jan', 'Feb', 'Mar', 'Q1 Revenue (3-D)', 'Q1 Profit (3-D)'])
for i, reg in enumerate(regions):
    r = 5 + i
    put(ws, f'A{r}', reg)
    put(ws, f'B{r}', f'=Jan!B{r}', 'link', USD)
    put(ws, f'C{r}', f'=Feb!B{r}', 'link', USD)
    put(ws, f'D{r}', f'=Mar!B{r}', 'link', USD)
    put(ws, f'E{r}', f'=SUM(Jan:Mar!B{r})', 'link', USD)
    put(ws, f'F{r}', f'=SUM(Jan:Mar!D{r})', 'link', USD)
put(ws, 'A9', 'Total', bold=True)
for col in 'BCDEF':
    put(ws, f'{col}9', f'=SUM({col}5:{col}8)', bold=True, fmt=USD)
put(ws, 'H5', '=Jan!B5', 'ftext')
put(ws, 'H6', '=SUM(Jan:Mar!B6)', 'ftext')
put(ws, 'H4', 'Formulas used', bold=True)
r = steps(ws, 12, [
    '1. Click E5 and press F2 to see the 3-D reference Jan:Mar!B5.',
    '2. Drag the Feb tab so it sits to the right of Mar, then back. Totals are unchanged while Feb stays between Jan and Mar.',
    '3. Group edit: click Jan, Shift+click Mar, type a note in F1. It appears on all three. Right-click > Ungroup Sheets.',
    '4. Insert a new sheet between Jan and Mar with the same layout: the 3-D totals include it automatically.',
], head='Try it')
r = steps(ws, r + 1, [
    'A link to another workbook looks like this (shown as text; red text is the convention for external links):',
], head='Linking to another workbook')
c = put(ws, f'A{r}', "='[Budget2026.xlsx]Jan'!$B$9", 'ftext'); c.font = font(color='FF0000')
r = steps(ws, r + 1, [
    'When the source is closed Excel shows the full path, e.g. =\'C:\\Reports\\[Budget2026.xlsx]Jan\'!$B$9',
    'Manage links: Data > Edit Links (Workbook Links) to update, change source or break a link.',
    'Data > Consolidate can also total ranges from several sheets or workbooks without writing formulas.',
])

# ---------------------------------------------------------------- 8 Auditing
ws = sheet('8 Auditing', '8. Formula auditing',
           'Guide topic 8. Use Formulas > Formula Auditing on the model, then work through the error table and the bug hunt.',
           [26, 14, 3, 30, 22, 3, 14, 10])
put(ws, 'A4', 'Mini income model', bold=True)
model = [('Units sold', 500, 'input', '#,##0'), ('Price per unit', 20, 'input', USD), ('Revenue', '=B5*B6', 'f', USD),
         ('COGS %', 0.6, 'input', '0%'), ('COGS', '=B7*B8', 'f', USD), ('Gross profit', '=B7-B9', 'f', USD),
         ('Operating expenses', 2500, 'input', USD), ('Operating income', '=B10-B11', 'f', USD),
         ('Tax rate', 0.25, 'input', '0%'), ('Net income', '=B12*(1-B13)', 'f', USD)]
for i, (lab, v, kind, fmt) in enumerate(model):
    r = 5 + i
    put(ws, f'A{r}', lab)
    put(ws, f'B{r}', v, 'input' if kind == 'input' else 'text', fmt)
put(ws, 'A15', 'All model inputs are invented.', 'note')
steps(ws, 4, [
    '1. Click B14 > Formulas > Trace Precedents. Click again to go a level deeper.',
    '2. Click B5 > Trace Dependents to see everything Units feeds.',
    '3. Remove Arrows clears them.',
    '4. Ctrl + ` (grave accent) toggles Show Formulas.',
    '5. Click B14 > Evaluate Formula > Evaluate to step through the calculation.',
    '6. Watch Window > Add Watch > B14, then change B5 from another sheet.',
], col=4, head='Auditing tools to try')
put(ws, 'G4', 'Test cells', bold=True)
put(ws, 'G5', 'Zero cell'); put(ws, 'H5', 0, 'input')
put(ws, 'G6', 'Text cell'); put(ws, 'H6', 'abc', 'input')

put(ws, 'A18', 'Error table: type the broken formula in a spare cell, read the error, then compare with the fix', bold=True)
hdr(ws, 19, 1, ['Error (shown with a leading #)', 'Fixed (live)', '', 'Broken formula to type', 'What causes it'])
errs = [
    ('DIV/0!', '=IF(H5=0,"n/a",B7/H5)', '=B7/H5', 'Dividing by zero or by an empty cell'),
    ('VALUE!', '=B7+N(H6)', '=B7+H6', 'Maths on text; N() turns text into 0'),
    ('N/A', '=IFERROR(VLOOKUP("Tablet",Data!$K$5:$M$9,3,FALSE),"Not found")',
     '=VLOOKUP("Tablet",Data!K5:M9,3,FALSE)', 'Lookup value is not in the list'),
    ('NAME?', '=SUM(B5:B6)', '=SUMM(B5:B6)', 'Misspelled function or undefined name'),
    ('REF!', None, '=B7 then delete row 7 (then Undo!)', 'The cell a formula pointed to was deleted'),
    ('#####', None, 'Narrow column B until numbers vanish', 'Column too narrow: not really an error'),
    ('Circular', None, 'Type =B14+1 into B14 (then Undo!)', 'A formula that refers to its own cell'),
]
for i, (e, fix, broken, why) in enumerate(errs):
    r = 20 + i
    put(ws, f'A{r}', e, bold=True).data_type = 's'
    if fix:
        put(ws, f'B{r}', fix, 'check', USD if e == 'VALUE!' else None, align='right')
    put(ws, f'D{r}', broken, 'ftext' if broken.startswith('=') and 'then' not in broken else 'text').data_type = 's'
    put(ws, f'E{r}', why, 'note')

put(ws, 'A29', 'Bug hunt: no error value, but the total is wrong. Why?', bold=True)
hdr(ws, 30, 1, ['Expense', 'Amount'])
for i, (lab, v) in enumerate([('Rent', 1200), ('Utilities', 800), ('Supplies', 450), ('Insurance', 300)]):
    put(ws, f'A{31 + i}', lab)
    put(ws, f'B{31 + i}', v, 'input', USD)
put(ws, 'A35', 'Total (deliberately wrong)', bold=True)
c = put(ws, 'B35', '=SUM(B31:B33)', fmt=USD, bold=True)
put(ws, 'A36', 'Correct total', 'check')
put(ws, 'B36', '=SUM(B31:B34)', 'check', USD)
put(ws, 'D31', 'Click B35 > Trace Precedents: the blue box stops at row 33.', 'note')
put(ws, 'D32', 'Excel also flags it with a green triangle: "Formula omits adjacent cells".', 'note')

# ---------------------------------------------------------------- 9 Input Control
ws = sheet('9 Input Control', '9. Controlling user input',
           'Guide topic 9. Every yellow cell has a data validation rule. Try to break each one.',
           [20, 16, 3, 16, 46, 34])
put(ws, 'A4', 'Order form', bold=True)
form = [('Product', 'Laptop', None), ('Quantity', 3, '0'), ('Order date', dt.date(2026, 10, 8), DATE),
        ('Discount', 0.1, '0%'), ('Customer code', 'C1234', None), ('Rush order?', 'No', None)]
for i, (lab, v, fmt) in enumerate(form):
    put(ws, f'A{5 + i}', lab)
    c = put(ws, f'B{5 + i}', v, 'input', fmt, align='right')
    from openpyxl.styles import Protection; c.protection = Protection(locked=False)
put(ws, 'A12', 'Unit price'); put(ws, 'B12', '=VLOOKUP(B5,Data!$K$5:$M$9,3,FALSE)', 'link', USD)
put(ws, 'A13', 'Order total', bold=True); put(ws, 'B13', '=B6*B12*(1-B8)', bold=True, fmt=USD)

def add_dv(cell, title, prompt, etitle, emsg, style='stop', **k):
    d = DataValidation(allow_blank=False, showInputMessage=True, showErrorMessage=True,
                       promptTitle=title, prompt=prompt, errorTitle=etitle, error=emsg, errorStyle=style, **k)
    ws.add_data_validation(d); d.add(cell)

add_dv('B5', 'Product', 'Pick a product from the list.', 'Not on the list', 'Choose one of the five products.',
       type='list', formula1='=Data!$K$5:$K$9')
add_dv('B6', 'Quantity', 'Whole number from 1 to 100.', 'Invalid quantity', 'Enter a whole number between 1 and 100.',
       type='whole', operator='between', formula1='1', formula2='100')
add_dv('B7', 'Order date', 'On or after Jan 1, 2026.', 'Invalid date', 'The date must be Jan 1, 2026 or later.',
       type='date', operator='greaterThanOrEqual', formula1='DATE(2026,1,1)')
add_dv('B8', 'Discount', '0% to 20%.', 'Large discount', 'Discounts above 20% need approval. Continue?',
       style='warning', type='decimal', operator='between', formula1='0', formula2='0.2')
add_dv('B9', 'Customer code', 'C followed by 4 characters, e.g. C1234.', 'Invalid code',
       'The code must be 5 characters and start with C.', type='custom', formula1='AND(LEN(B9)=5,LEFT(B9,1)="C")')
add_dv('B10', 'Rush', 'Yes or No.', 'Invalid', 'Choose Yes or No.', type='list', formula1='"Yes,No"')

hdr(ws, 4, 4, ['Cell', 'Rule (Data > Data Validation)', 'Try typing this'])
rules = [('B5 Product', 'List, source =Data!$K$5:$K$9', 'Tablet (rejected)'),
         ('B6 Quantity', 'Whole number between 1 and 100', '0, 2.5 or 500 (all rejected)'),
         ('B7 Order date', 'Date >= DATE(2026,1,1)', '12/31/2025 (rejected)'),
         ('B8 Discount', 'Decimal 0 to 0.2, Warning style', '0.5 (warning: you may continue)'),
         ('B9 Customer code', 'Custom =AND(LEN(B9)=5,LEFT(B9,1)="C")', 'X1234 or C12 (rejected)'),
         ('B10 Rush order?', 'List typed in the box: Yes,No', 'Maybe (rejected)')]
for i, row in enumerate(rules):
    for j, v in enumerate(row):
        put(ws, f'{L(4 + j)}{5 + i}', v)
steps(ws, 16, [
    '1. The six yellow cells are already unlocked (Format Cells > Protection > Locked is unticked).',
    '2. Review > Protect Sheet > OK (a password is optional).',
    '3. Now try to type over B13: Excel refuses. The yellow cells still accept input.',
    '4. Review > Unprotect Sheet to switch it off.',
    '5. Data > Data Validation > Circle Invalid Data highlights entries that broke a rule before it existed.',
], head='Protect the sheet so only the inputs can change')

# ---------------------------------------------------------------- 10 Financial
ws = sheet('10 Financial', '10. Financial functions',
           'Guide topic 10. Money you pay out is negative, money you receive is positive. Rate and periods must use the same time unit.',
           [34, 14, 34, 3, 3, 8, 13, 12, 12, 12, 13])
put(ws, 'A4', 'Loan amount'); put(ws, 'B4', 25000, 'input', USD)
put(ws, 'A5', 'Annual interest rate'); put(ws, 'B5', 0.06, 'input', '0.00%')
put(ws, 'A6', 'Term (years)'); put(ws, 'B6', 5, 'input')
put(ws, 'A7', 'Payments per year'); put(ws, 'B7', 12, 'input')
put(ws, 'C4', 'Loan inputs are invented', 'note')
fin = [
    (9, 'Monthly payment', '=PMT(B5/B7,B6*B7,-B4)', USD2),
    (10, 'Total paid', '=B9*B6*B7', USD2),
    (11, 'Total interest', '=B10-B4', USD2),
    (12, 'Interest in payment 1', '=IPMT(B5/B7,1,B6*B7,-B4)', USD2),
    (13, 'Principal in payment 1', '=PPMT(B5/B7,1,B6*B7,-B4)', USD2),
    (14, 'Check: number of payments', '=NPER(B5/B7,-B9,B4)', '0.0'),
    (15, 'Check: annual rate', '=RATE(B6*B7,-B9,B4)*B7', '0.00%'),
    (18, 'Monthly deposit', 200, USD), (19, 'Annual return', 0.05, '0.00%'), (20, 'Years', 10, '0'),
    (21, 'Future value of the savings', '=FV(B19/12,B20*12,-B18)', USD2),
    (24, 'Amount received in the future', 10000, USD), (25, 'Years until received', 5, '0'),
    (26, 'Discount rate', 0.05, '0.00%'),
    (27, 'Present value today', '=PV(B26,B25,0,-B24)', USD2),
    (30, 'Discount rate', 0.08, '0.00%'), (31, 'Year 0 (investment)', -10000, USD),
    (32, 'Year 1', 3000, USD), (33, 'Year 2', 4000, USD), (34, 'Year 3', 4000, USD), (35, 'Year 4', 3000, USD),
    (36, 'Net present value', '=NPV(B30,B32:B35)+B31', USD2),
    (37, 'Internal rate of return', '=IRR(B31:B35)', '0.00%'),
    (40, 'Asset cost', 10000, USD), (41, 'Salvage value', 1000, USD), (42, 'Useful life (years)', 5, '0'),
    (43, 'Straight-line depreciation per year', '=SLN(B40,B41,B42)', USD2),
]
for r, lab, v, fmt in fin:
    put(ws, f'A{r}', lab)
    if isinstance(v, str):
        put(ws, f'B{r}', v, fmt=fmt)
        put(ws, f'C{r}', v, 'ftext')
    else:
        put(ws, f'B{r}', v, 'input', fmt)
for r, t_ in [(8, 'Loan results'), (17, 'Saving for a goal (FV)'), (23, 'Value today of future money (PV)'),
              (29, 'Evaluating a project (NPV and IRR)'), (39, 'Depreciation (SLN)')]:
    put(ws, f'A{r}', t_, bold=True)
put(ws, 'C36', '=NPV(B30,B32:B35)+B31   NPV starts at year 1, so add year 0 separately', 'ftext')

put(ws, 'F3', 'Amortization schedule', bold=True)
hdr(ws, 4, 6, ['Period', 'Beginning', 'Payment', 'Interest', 'Principal', 'Ending'])
for n in range(1, 61):
    r = 4 + n
    put(ws, f'F{r}', n)
    put(ws, f'G{r}', '=$B$4' if n == 1 else f'=K{r - 1}', fmt=USD2)
    put(ws, f'H{r}', '=$B$9', fmt=USD2)
    put(ws, f'I{r}', f'=G{r}*$B$5/$B$7', fmt=USD2)
    put(ws, f'J{r}', f'=H{r}-I{r}', fmt=USD2)
    put(ws, f'K{r}', f'=G{r}-J{r}', fmt=USD2)
put(ws, 'F66', 'Total', bold=True)
put(ws, 'I66', '=SUM(I5:I64)', bold=True, fmt=USD2)
put(ws, 'J66', '=SUM(J5:J64)', bold=True, fmt=USD2)
put(ws, 'F67', 'Schedule is built for the 60 payments of the default 5-year monthly loan.', 'note')

# ---------------------------------------------------------------- 11 Solver Linear
ws = sheet('11 Solver Linear', '11. Solver: linear optimization',
           'Guide topic 11. A furniture maker chooses how many tables and chairs to build to maximise profit. Numbers are invented.',
           [30, 12, 12, 12, 6, 12])
hdr(ws, 4, 2, ['Tables', 'Chairs'])
put(ws, 'A5', 'Units to make (changing cells)')
put(ws, 'B5', 0, 'todo'); put(ws, 'C5', 0, 'todo')
put(ws, 'A6', 'Profit per unit'); put(ws, 'B6', 50, 'input', USD); put(ws, 'C6', 30, 'input', USD)
put(ws, 'A8', 'Total profit (objective)', bold=True)
put(ws, 'B8', '=SUMPRODUCT(B5:C5,B6:C6)', bold=True, fmt=USD)
hdr(ws, 10, 1, ['Resource (hours per unit)', 'Tables', 'Chairs', 'Used', '', 'Available'])
for r, lab, a, b, avail in [(11, 'Carpentry hours', 4, 3, 240), (12, 'Finishing hours', 2, 1, 100)]:
    put(ws, f'A{r}', lab)
    put(ws, f'B{r}', a, 'input'); put(ws, f'C{r}', b, 'input')
    put(ws, f'D{r}', f'=SUMPRODUCT($B$5:$C$5,B{r}:C{r})')
    put(ws, f'E{r}', '<=', align='center')
    put(ws, f'F{r}', avail, 'input')
r = steps(ws, 15, [
    '0. One-time setup: File > Options > Add-ins > Manage: Excel Add-ins > Go > tick Solver Add-in.',
    '1. Data > Solver.',
    '2. Set Objective: $B$8     To: Max',
    '3. By Changing Variable Cells: $B$5:$C$5',
    '4. Add constraint: $D$11:$D$12 <= $F$11:$F$12',
    '5. Tick "Make Unconstrained Variables Non-Negative".',
    '6. Select a Solving Method: Simplex LP.',
    '7. Solve > Keep Solver Solution. Tick Answer and Sensitivity reports to see them on new sheets.',
], head='Solver setup')
put(ws, f'A{r + 1}', 'Answer check', bold=True)
put(ws, f'A{r + 2}', 'Tables', 'check'); put(ws, f'B{r + 2}', 30, 'check')
put(ws, f'A{r + 3}', 'Chairs', 'check'); put(ws, f'B{r + 3}', 40, 'check')
put(ws, f'A{r + 4}', 'Maximum profit', 'check')
put(ws, f'B{r + 4}', f'=B{r + 2}*B6+B{r + 3}*C6', 'check', USD)
put(ws, f'A{r + 5}', 'Both resources are fully used at the optimum (binding constraints).', 'note')

# ---------------------------------------------------------------- 12 Solver Nonlinear
ws = sheet('12 Solver Nonlinear', '12. Solver: nonlinear optimization',
           'Guide topic 12. Demand falls as price rises, so profit is a curve. Solver finds the top of it. Numbers are invented.',
           [32, 14, 3, 10, 12])
put(ws, 'A4', 'Price (changing cell)'); put(ws, 'B4', 10, 'todo', USD2)
put(ws, 'A5', 'Unit cost'); put(ws, 'B5', 6, 'input', USD2)
put(ws, 'A6', 'Fixed cost'); put(ws, 'B6', 2000, 'input', USD)
put(ws, 'A7', 'Demand at price 0'); put(ws, 'B7', 1200, 'input', '#,##0')
put(ws, 'A8', 'Units lost per $1 of price'); put(ws, 'B8', 40, 'input')
put(ws, 'A10', 'Demand (units)'); put(ws, 'B10', '=B7-B8*B4', fmt='#,##0')
put(ws, 'A11', 'Revenue'); put(ws, 'B11', '=B4*B10', fmt=USD)
put(ws, 'A12', 'Total cost'); put(ws, 'B12', '=B6+B5*B10', fmt=USD)
put(ws, 'A13', 'Profit (objective)', bold=True); put(ws, 'B13', '=B11-B12', bold=True, fmt=USD)
hdr(ws, 4, 4, ['Price', 'Profit'])
for i, p in enumerate([6, 10, 14, 18, 22, 26, 30]):
    r = 5 + i
    put(ws, f'D{r}', p, 'input', USD)
    put(ws, f'E{r}', f'=(D{r}-$B$5)*($B$7-$B$8*D{r})-$B$6', fmt=USD)
r = steps(ws, 16, [
    '1. Data > Solver.',
    '2. Set Objective: $B$13     To: Max',
    '3. By Changing Variable Cells: $B$4',
    '4. Add constraints: $B$4 >= $B$5   and   $B$4 <= 30',
    '5. Select a Solving Method: GRG Nonlinear (price times demand makes the model nonlinear).',
    '6. Solve. Then start again from price 28 and solve: a smooth single-peak curve gives the same answer.',
    '7. For bumpy models tick Options > GRG Nonlinear > Use Multistart, or try Evolutionary.',
], head='Solver setup')
put(ws, f'A{r + 1}', 'Answer check', bold=True)
put(ws, f'A{r + 2}', 'Best price', 'check'); put(ws, f'B{r + 2}', '=(B7/B8+B5)/2', 'check', USD2)
put(ws, f'A{r + 3}', 'Profit at best price', 'check')
put(ws, f'B{r + 3}', f'=(B{r + 2}-B5)*(B7-B8*B{r + 2})-B6', 'check', USD)
put(ws, f'A{r + 4}', 'Calculus check: profit peaks halfway between unit cost and the price where demand hits zero.', 'note')
sc = ScatterChart(); sc.title = 'Profit curve'; sc.style = 13
sc.x_axis.title = 'Price ($)'; sc.y_axis.title = 'Profit ($)'
sc.x_axis.delete = False; sc.y_axis.delete = False; sc.legend = None
s = Series(Reference(ws, min_col=5, min_row=5, max_row=11), Reference(ws, min_col=4, min_row=5, max_row=11))
s.smooth = True
sc.series.append(s); sc.height = 7.5; sc.width = 13
ws.add_chart(sc, 'G4')

# ---------------------------------------------------------------- 13 Data Tables
ws = sheet('13 Data Tables', '13. One- and two-variable data tables',
           'Guide topic 13. A data table reruns one formula for a list of input values. Fill the yellow areas with Data > What-If Analysis > Data Table.',
           [22, 13, 13, 13, 3, 13, 13, 13])
put(ws, 'A4', 'Loan amount'); put(ws, 'B4', 25000, 'input', USD)
put(ws, 'A5', 'Annual rate'); put(ws, 'B5', 0.06, 'input', '0.00%')
put(ws, 'A6', 'Years'); put(ws, 'B6', 5, 'input')
put(ws, 'A7', 'Monthly payment', bold=True); put(ws, 'B7', '=PMT(B5/12,B6*12,-B4)', bold=True, fmt=USD2)
put(ws, 'A9', 'One-variable table: payment at different rates', bold=True)
put(ws, 'A10', 'Rate'); put(ws, 'B10', '=B7', fmt=USD2)
put(ws, 'D9', 'Answer check', bold=True)
for i, rate in enumerate([0.04, 0.05, 0.06, 0.07, 0.08]):
    r = 11 + i
    put(ws, f'A{r}', rate, 'input', '0.00%')
    put(ws, f'B{r}', None, 'todo', USD2)
    put(ws, f'D{r}', f'=PMT(A{r}/12,$B$6*12,-$B$4)', 'check', USD2)
put(ws, 'F10', '1. Select A10:B15', 'text')
put(ws, 'F11', '2. Data > What-If Analysis > Data Table')
put(ws, 'F12', '3. Leave Row input cell empty')
put(ws, 'F13', '4. Column input cell: $B$5  > OK')
put(ws, 'F14', 'B11:B15 should equal the green column.', 'note')

put(ws, 'A18', 'Two-variable table: payment by rate (down) and years (across)', bold=True)
put(ws, 'A19', '=B7', fmt=USD2)
put(ws, 'F18', 'Answer check', bold=True)
for j, y in enumerate([3, 4, 5]):
    put(ws, f'{L(2 + j)}19', y, 'input')
    put(ws, f'{L(6 + j)}19', f'={L(2 + j)}19', 'check')
for i, rate in enumerate([0.04, 0.05, 0.06, 0.07, 0.08]):
    r = 20 + i
    put(ws, f'A{r}', rate, 'input', '0.00%')
    for j in range(3):
        put(ws, f'{L(2 + j)}{r}', None, 'todo', USD2)
        put(ws, f'{L(6 + j)}{r}', f'=PMT($A{r}/12,{L(6 + j)}$19*12,-$B$4)', 'check', USD2)
steps(ws, 27, [
    '1. Select A19:D24 (the formula sits in the top-left corner).',
    '2. Data > What-If Analysis > Data Table.',
    '3. Row input cell: $B$6 (the years run across the top row).',
    '4. Column input cell: $B$5 (the rates run down the left column) > OK.',
    '5. Click B20: the formula bar shows {=TABLE(B6,B5)}. You cannot edit one result cell on its own.',
], head='Two-variable steps')

# ---------------------------------------------------------------- 14 Scenarios
ws = sheet('14 Scenarios', '14. Scenario management',
           'Guide topic 14. Save sets of input values as named scenarios and switch between them. Numbers are invented.',
           [24, 14, 3, 14, 10, 10, 11, 16])
put(ws, 'A4', 'Units sold'); put(ws, 'B4', 1000, 'input', '#,##0')
put(ws, 'A5', 'Price'); put(ws, 'B5', 50, 'input', USD)
put(ws, 'A6', 'Unit cost'); put(ws, 'B6', 30, 'input', USD)
put(ws, 'A7', 'Fixed costs'); put(ws, 'B7', 15000, 'input', USD)
put(ws, 'A9', 'Revenue'); put(ws, 'B9', '=B4*B5', fmt=USD)
put(ws, 'A10', 'Total cost'); put(ws, 'B10', '=B7+B4*B6', fmt=USD)
put(ws, 'A11', 'Profit', bold=True); put(ws, 'B11', '=B9-B10', bold=True, fmt='$#,##0;($#,##0);-')
hdr(ws, 4, 4, ['Scenario', 'Units', 'Price', 'Unit cost', 'Expected profit'])
for i, (nm, u, p, c_) in enumerate([('Best', 1200, 52, 28), ('Base', 1000, 50, 30), ('Worst', 800, 48, 32)]):
    r = 5 + i
    put(ws, f'D{r}', nm)
    put(ws, f'E{r}', u, 'input', '#,##0'); put(ws, f'F{r}', p, 'input', USD); put(ws, f'G{r}', c_, 'input', USD)
    put(ws, f'H{r}', f'=E{r}*(F{r}-G{r})-$B$7', 'check', '$#,##0;($#,##0);-')
r = steps(ws, 14, [
    '1. Data > What-If Analysis > Scenario Manager > Add.',
    '2. Scenario name: Best     Changing cells: $B$4:$B$6 > OK.',
    '3. Enter 1200, 52 and 28 > Add. Repeat for Base and Worst using the table above > OK.',
    '4. Pick a scenario > Show. Profit in B11 should equal the green check for that scenario.',
    '5. Summary > Result cells: $B$11 > OK. Excel builds a Scenario Summary sheet comparing all three.',
    '6. Tip: name the cells first (Formulas > Define Name) so the summary shows names, not addresses.',
], head='Scenario Manager steps')
r = steps(ws, r + 1, [
    '1. Show the Base scenario (1000, 50, 30).',
    '2. Data > What-If Analysis > Goal Seek.',
    '3. Set cell: $B$11    To value: 0    By changing cell: $B$4 > OK.',
], head='Goal Seek: how many units to break even?')
put(ws, f'A{r}', 'Answer check: break-even units', 'check')
put(ws, f'B{r}', '=E6', 'text')
ws[f'B{r}'].value = '=B7/(F6-G6)'
ws[f'B{r}'].fill = CHECK_FILL; ws[f'B{r}'].border = BORDER; ws[f'B{r}'].number_format = '#,##0'

# ---------------------------------------------------------------- 15 Macros
ws = sheet('15 Macros', '15. Excel macros',
           'Guide topic 15. First: File > Save As > Excel Macro-Enabled Workbook (*.xlsm). An .xlsx file discards macros when saved.',
           [70, 3, 12, 12])
hdr(ws, 4, 3, ['Order', 'Revenue'])
for i in range(5):
    r = 5 + i
    put(ws, f'C{r}', i + 1)
    put(ws, f'D{r}', f'=Data!H{5 + i}', 'link', USD)
r = steps(ws, 4, [
    '0. Show the Developer tab: File > Options > Customize Ribbon > tick Developer.',
    '1. Select C4:D4. Developer > Record Macro. Name: FormatHeader, Shortcut: Ctrl+Shift+H, Store in: This Workbook.',
    '2. Make the cells bold, fill them green, centre them. Developer > Stop Recording.',
    '3. Select any other cells and press Ctrl+Shift+H to replay it.',
    '4. Developer > Macros > FormatHeader > Edit opens the VBA editor (Alt+F11) with the code Excel wrote.',
    '5. Use Relative References (toggle before recording) makes a macro act from the active cell instead of fixed addresses.',
], head='Record your first macro')
r = steps(ws, r + 1, [
    'Alt+F11 > Insert > Module, then type or paste the code below. Run with F5 or Developer > Macros.',
], head='Three short procedures to paste')
code = '''Sub FormatHeader()
    ' Formats whatever cells are selected as a header row
    With Selection
        .Font.Bold = True
        .Font.Color = vbWhite
        .Interior.Color = RGB(33, 115, 70)
        .HorizontalAlignment = xlCenter
    End With
End Sub

Sub AddTimestamp()
    ' Stamps the active cell with the current date and time
    ActiveCell.Value = Now
    ActiveCell.NumberFormat = "yyyy-mm-dd hh:mm"
End Sub

Sub HighlightBigOrders()
    ' Loops through the revenue cells and colours any over 2,000
    Dim c As Range
    For Each c In Worksheets("15 Macros").Range("D5:D9")
        If c.Value > 2000 Then c.Interior.Color = vbYellow
    Next c
End Sub'''
for ln in code.split('\n'):
    put(ws, f'A{r}', ln if ln else None, 'code')
    r += 1
r = steps(ws, r + 1, [
    'HighlightBigOrders should colour D5 and D6 (the two orders above $2,000).',
    'Assign a macro to a button: Developer > Insert > Button (Form Control), draw it, pick the macro.',
    'Security: File > Options > Trust Center > Macro Settings. Only enable macros in files you trust.',
], head='After running')

wb.save('Advanced_Excel_Practice_Workbook.xlsx')

with open('regional_targets.csv', 'w') as f:
    f.write('Region,Quarter,Target\nEast,Q1,7000\nWest,Q1,9000\nNorth,Q1,6000\nSouth,Q1,5500\n')
print('built')
